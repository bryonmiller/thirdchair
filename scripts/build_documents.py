"""Build the proposal, one-page brief, and submission copy as Word files."""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INK = "151B22"
BLUE = "2E74B5"
GREEN = "1C5961"
CREAM = "F5F0E4"
SAND = "D8D0C0"
PALE = "F4F6F9"
GRAY = "596067"
WHITE = "FFFFFF"
RED = "A9462F"
AMBER = "E5AD46"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, *, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, *, color: str = SAND, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(GRAY)


def add_hyperlink(paragraph, text: str, url: str, *, color: str = BLUE) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color_node, underline])
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://\S+)")


def add_inline(paragraph, text: str, *, size: float | None = None) -> None:
    position = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if size:
                run.font.size = Pt(size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = rgb(BLUE)
            if size:
                run.font.size = Pt(size)
        else:
            url = token.rstrip(".,;)")
            add_hyperlink(paragraph, url, url)
            if len(url) < len(token):
                paragraph.add_run(token[len(url) :])
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        if size:
            run.font.size = Pt(size)


def setup_document(doc: Document, *, compact: bool = False, margins: float = 1.0) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(margins)
    section.bottom_margin = Inches(margins)
    section.left_margin = Inches(margins)
    section.right_margin = Inches(margins)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.1 if compact else 11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(3 if compact else 8)
    normal.paragraph_format.line_spacing = 1.08 if compact else 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if compact else WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Title": (28 if not compact else 22, INK, 0, 10),
        "Subtitle": (13 if not compact else 11, GRAY, 0, 10),
        "Heading 1": (16 if not compact else 13, BLUE, 16, 8),
        "Heading 2": (13 if not compact else 11, BLUE, 12, 6),
        "Heading 3": (12 if not compact else 10, GREEN, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(9.1 if compact else 10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(2 if compact else 4)
        style.paragraph_format.line_spacing = 1.08 if compact else 1.2


def add_header_footer(doc: Document, document_label: str, *, compact: bool = False) -> None:
    section = doc.sections[0]
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5 if not compact else 7.2))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.4)
    table.columns[1].width = Inches(2.1 if not compact else 2.8)
    for cell in table.rows[0].cells:
        cell_margins(cell, top=0, start=0, bottom=35, end=0)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run("THE THIRD CHAIR")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(7.5)
    run.font.color.rgb = rgb(BLUE)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    run = right.add_run(document_label.upper())
    run.font.name = "Calibri"
    run.font.size = Pt(7.5)
    run.font.color.rgb = rgb(GRAY)

    footer = section.footer
    table = footer.add_table(rows=1, cols=3, width=Inches(6.5 if not compact else 7.2))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in table.rows[0].cells:
        cell_margins(cell, top=30, start=0, bottom=0, end=0)
    entries = ("BRYON MILLER", "CHINATALK EVALS FOR THE SITUATION ROOM", "")
    for index, value in enumerate(entries):
        paragraph = table.cell(0, index).paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        if index == 1:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if index == 2:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            page_number(paragraph)
        else:
            run = paragraph.add_run(value)
            run.font.name = "Calibri"
            run.font.size = Pt(7.5)
            run.font.color.rgb = rgb(GRAY)


def add_rule(paragraph, *, color: str = GREEN, size: int = 14) -> None:
    props = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    props.append(borders)


def add_code_block(doc: Document, lines: Iterable[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, PALE)
    set_cell_border(cell, color=SAND, size=4)
    cell_margins(cell, top=100, start=140, bottom=100, end=140)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.2)
        run.font.color.rgb = rgb(INK)


def render_markdown(doc: Document, source: str) -> None:
    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    in_code = False

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0)
        paragraph.paragraph_format.right_indent = Inches(0)
        paragraph.paragraph_format.first_line_indent = Inches(0)
        add_inline(paragraph, " ".join(line.strip() for line in paragraph_lines))
        paragraph_lines.clear()

    for raw_line in source.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(doc, code_lines)
                code_lines.clear()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            flush_paragraph()
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)), 3)
            doc.add_heading(heading.group(2), level=level)
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.keep_together = True
            add_inline(paragraph, bullet.group(1))
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.right_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            number = paragraph.add_run(f"{numbered.group(1)}.  ")
            number.bold = True
            add_inline(paragraph, numbered.group(2))
            continue
        paragraph_lines.append(line)
    flush_paragraph()
    if code_lines:
        add_code_block(doc, code_lines)


def proposal_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)
    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(12)
    run = label.add_run("CHINATALK EVALS FOR THE SITUATION ROOM")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(GREEN)
    add_rule(label, color=GREEN, size=18)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(8)
    title.add_run("The Third Chair")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    subtitle.paragraph_format.left_indent = Inches(0)
    subtitle.paragraph_format.right_indent = Inches(0)
    subtitle.paragraph_format.first_line_indent = Inches(0)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("Can an AI stop a cyber crisis it might also help start?")
    run.font.name = "Calibri"
    run.font.size = Pt(17)
    run.font.color.rgb = rgb(BLUE)

    question = doc.add_table(rows=1, cols=1)
    question.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = question.cell(0, 0)
    shade(cell, CREAM)
    set_cell_border(cell, color=SAND, size=8)
    cell_margins(cell, top=180, start=200, bottom=180, end=200)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run("THE QUESTION")
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(GREEN)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        "Two models play rival governments. A third gets two chances to lower escalation without rewarding a lie, then all three swap seats."
    )
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = rgb(INK)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(26)
    meta.paragraph_format.space_after = Pt(0)
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = meta.add_run("Bryon Miller")
    run.bold = True
    run.font.size = Pt(12)
    meta.add_run("\nAugust 2026 | Version 0.3.0").font.color.rgb = rgb(GRAY)
    status = doc.add_paragraph()
    status.paragraph_format.space_before = Pt(18)
    status.paragraph_format.space_after = Pt(0)
    status.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = status.add_run("STATUS  ")
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(RED)
    run = status.add_run("Tournament engine complete. Comparative model pilot awaiting API access.")
    run.font.size = Pt(9.5)
    run.font.color.rgb = rgb(GRAY)


def build_proposal(output: Path) -> None:
    doc = Document()
    setup_document(doc, compact=False, margins=1.0)
    add_header_footer(doc, "Proposal")
    proposal_cover(doc)
    doc.add_page_break()

    source = (ROOT / "docs" / "proposal.md").read_text(encoding="utf-8")
    body = source[source.index("## Summary") :]
    render_markdown(doc, body)
    doc.core_properties.title = "The Third Chair: evaluation proposal"
    doc.core_properties.subject = "ChinaTalk Evals for the Situation Room"
    doc.core_properties.author = "Bryon Miller"
    doc.core_properties.keywords = "AI evaluation, cyber crisis, mediation, multi-agent tournament"
    doc.save(output)


def source_sections(path: Path) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    current = ""
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        if raw_line.startswith("## "):
            current = raw_line[3:].strip()
            sections[current] = []
        elif current and raw_line.strip():
            sections[current].append(raw_line.strip())
    return sections


def cell_section(cell, title: str, lines: list[str]) -> None:
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(GREEN)
    for line in lines:
        if line.startswith("- "):
            paragraph = cell.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1)
            add_inline(paragraph, line[2:], size=8.7)
        else:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, line, size=8.7)


def build_brief(output: Path) -> None:
    doc = Document()
    setup_document(doc, compact=True, margins=0.55)
    add_header_footer(doc, "One-page brief", compact=True)
    sections = source_sections(ROOT / "docs" / "brief.md")

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(2)
    title.add_run("The Third Chair")
    deck = doc.add_paragraph()
    deck.paragraph_format.space_after = Pt(8)
    run = deck.add_run("A rotating model tournament for truth-preserving cyber-crisis mediation")
    run.font.size = Pt(11)
    run.font.color.rgb = rgb(BLUE)

    question_line = sections["One-page evaluation brief"][0]
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    shade(cell, CREAM)
    set_cell_border(cell, color=SAND, size=6)
    cell_margins(cell, top=85, start=120, bottom=85, end=120)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, question_line, size=9.2)

    columns = doc.add_table(rows=1, cols=2)
    columns.alignment = WD_TABLE_ALIGNMENT.CENTER
    columns.autofit = False
    columns.columns[0].width = Inches(3.52)
    columns.columns[1].width = Inches(3.52)
    left, right = columns.rows[0].cells
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell_margins(cell, top=90, start=75, bottom=40, end=75)
    cell_section(left, "What happens", sections["What happens in a match"])
    cell_section(left, "The hard part", sections["The hard part"])
    cell_section(left, "Rotation and control", sections["Rotation and control"])
    cell_section(right, "What the pilot reports", sections["What the pilot reports"])
    cell_section(right, "Pilot size and status", sections["Pilot size and status"])
    cell_section(right, "Why Bryon", sections["Why Bryon"])

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(0)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "The board stays empty until every result can be traced to a stored response and exact model identifier."
    )
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(RED)

    doc.core_properties.title = "The Third Chair: one-page brief"
    doc.core_properties.author = "Bryon Miller"
    doc.save(output)


def build_submission(output: Path) -> None:
    doc = Document()
    setup_document(doc, compact=False, margins=0.8)
    add_header_footer(doc, "Submission copy")
    source = source_sections(ROOT / "docs" / "submission_copy.md")

    title = doc.add_paragraph(style="Title")
    title.add_run("ChinaTalk submission copy")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("The Third Chair | Bryon Miller")
    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(12)
    add_inline(
        note,
        "Prepared for direct entry into the contest form. Confirm the two intended public URLs before submission.",
    )

    order = [
        "Name",
        "Email",
        "LinkedIn",
        "Two-sentence author bio",
        "Evaluation title",
        "Abstract",
        "Microsite",
        "GitHub repository",
        "Interest in working for ChinaTalk",
        "Submission note",
    ]
    for heading in order:
        lines = source[heading]
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(1.65)
        table.columns[1].width = Inches(5.1)
        keep_row_together(table.rows[0])
        label_cell, value_cell = table.rows[0].cells
        shade(label_cell, PALE)
        shade(value_cell, WHITE)
        for cell in (label_cell, value_cell):
            set_cell_border(cell, color=SAND, size=5)
            cell_margins(cell, top=100, start=120, bottom=100, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        label = label_cell.paragraphs[0]
        label.paragraph_format.space_after = Pt(0)
        run = label.add_run(heading.upper())
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = rgb(GREEN)
        value = value_cell.paragraphs[0]
        value.paragraph_format.space_after = Pt(0)
        add_inline(value, " ".join(lines))
        if heading == "Abstract":
            count = len(" ".join(lines).split())
            count_run = label_cell.add_paragraph().add_run(f"{count} words")
            count_run.font.size = Pt(8)
            count_run.font.color.rgb = rgb(GRAY)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        spacer.paragraph_format.line_spacing = Pt(2)

    checklist = doc.add_table(rows=1, cols=1)
    checklist.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = checklist.cell(0, 0)
    shade(cell, CREAM)
    set_cell_border(cell, color=SAND, size=6)
    cell_margins(cell, top=130, start=150, bottom=130, end=150)
    heading = cell.paragraphs[0]
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run("RELEASE CHECKLIST")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(GREEN)
    intro = cell.add_paragraph()
    intro.paragraph_format.space_after = Pt(5)
    intro.add_run("Keep this checklist out of the form response:").italic = True
    for item in (
        "Run the balanced pilot and inspect all 18 match tapes.",
        "Confirm the repository and microsite in a private browser window.",
        "Check all four downloads and the source archive checksum.",
        "Replace every draft URL note with the verified public address.",
        "Keep the model-result caveat if the pilot is not complete.",
        "Submit by September 1, 2026.",
    ):
        paragraph = cell.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)

    warning = doc.add_paragraph()
    warning.paragraph_format.space_before = Pt(8)
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = warning.add_run(
        "The evaluation package is complete, but the strongest submission will include traceable pilot results."
    )
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(RED)

    doc.core_properties.title = "The Third Chair: contest submission copy"
    doc.core_properties.author = "Bryon Miller"
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "deliverables")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    build_proposal(args.output_dir / "third-chair-proposal.docx")
    build_brief(args.output_dir / "third-chair-brief.docx")
    build_submission(args.output_dir / "third-chair-submission.docx")
    for path in sorted(args.output_dir.glob("*.docx")):
        print(path)


if __name__ == "__main__":
    main()
